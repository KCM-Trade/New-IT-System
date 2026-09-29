"""xlsx export of one exec-compensation query (03 §4.3 /export, 03 §5).

Two sheets, 汇总 (summary) and 明细 (every in-range fill, view=all). Cell A1
of both sheets carries the three mandatory notices (01 D13 / D14 / D17) and
the basis on one line (it overflows into the empty cells to its right), so a
sheet forwarded on its own still says what it is. Numbers are written as
numbers; nothing is recomputed here. Write-only workbook: a 100k-row detail
sheet streams instead of building a cell object graph. Detail rows are the
query core's plain row dicts (``query.export_rows``), not pydantic objects —
at the 150k deal cap that difference is hundreds of MB.
"""

from __future__ import annotations

import datetime as dt
import io
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from typing import Iterable, Mapping

from app.schemas.exec_compensation import SummaryResponse

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

BASIS_LINE = "口径 = 客户下单时的请求价；只统计客户主动市价单；金额以正负相抵为准（USD，CEN 已 ÷100）"

ROUNDING_NOTE = "逐笔金额保留 4 位小数，合计按未舍入值计算，逐笔相加可能与合计有 ±0.01 以内的差异"


def notices(s: SummaryResponse) -> str:
    ex = s.data.excluded_open
    as_of = s.as_of.isoformat()
    return (
        "①只適用於市價單開倉和平倉的差價　"
        f"②数据截至 {as_of}（查询日前一天，MT 服务器日），当天的成交不计入　"
        f"③截至 {as_of} 仍未完全平仓的仓位不计算：本次共剔除 {ex.positions} 个仓位、"
        f"{ex.deals} 笔成交、{ex.lots:g} 手"
    )


def _bold(ws, value, color: str | None = None) -> WriteOnlyCell:
    c = WriteOnlyCell(ws, value=value)
    c.font = Font(bold=True, color=color)
    return c


def _head(ws, s: SummaryResponse) -> None:
    ws.append([_bold(ws, notices(s) + "　" + BASIS_LINE, "C00000")])
    ws.append([])


def _subject_label(s: SummaryResponse) -> str:
    subj = s.data.subject
    if subj.client_id is not None:
        return f"client {subj.client_id}"
    return f"account {subj.login_sid}"


def filename(s: SummaryResponse) -> str:
    subj = s.data.subject
    who = f"client{subj.client_id}" if subj.client_id is not None else f"login{subj.login_sid}"
    q = s.data.query
    return f"exec-compensation_{who}_{q.date_from}_{q.date_to}_asof{s.as_of}.xlsx"


def content_disposition(name: str) -> str:
    ascii_name = name.encode("ascii", "ignore").decode("ascii") or "export.xlsx"
    return f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(name)}'


def _summary_sheet(ws, s: SummaryResponse) -> None:
    d, q, cov = s.data, s.data.query, s.coverage
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 22
    _head(ws, s)
    kv = [
        ("主体", _subject_label(s)),
        ("MT5 账户", ", ".join(f"5-{x}" for x in d.subject.logins)),
        ("日期范围（MT 服务器日）", f"{q.date_from} ~ {q.date_to}"
         + ("（已截到 as_of）" if q.date_to_clipped else "")),
        ("数据截止 as_of", s.as_of.isoformat()),
        ("数据完整", "是" if cov.complete else f"否：{cov.reason or ''}"),
        ("未找到订单行的成交（不计入）", cov.unmatched_deals),
        ("", ""),
        ("补偿（正负相抵，USD）", d.comp_net_usd),
        ("补偿（只补正数，USD）", d.comp_positive_usd),
        ("成交笔数", d.deals),
        ("手数", d.lots),
        ("更差 / 不变 / 更好", f"{d.outcomes.worse} / {d.outcomes.same} / {d.outcomes.better}"),
        ("延迟中位 / p95 / 最大 / 最小（ms）",
         f"{d.delay.median_ms} / {d.delay.p95_ms} / {d.delay.max_ms} / {d.delay.min_ms}"),
        ("单笔最大补偿（USD）", d.max_single_comp_usd),
        ("", ""),
        ("未完全平仓剔除：仓位数", d.excluded_open.positions),
        ("未完全平仓剔除：成交笔数", d.excluded_open.deals),
        ("未完全平仓剔除：手数", d.excluded_open.lots),
        ("未完全平仓剔除：若计入金额（USD）", d.excluded_open.comp_net_usd_if_counted),
        ("异常仓位数（不计入）", d.anomaly_positions),
        ("口径版本 calc_version", s.basis.calc_version),
    ]
    for k, v in kv:
        ws.append([k, v])
    ws.append([])
    ws.append([ROUNDING_NOTE])

    def table(title: str, rows, key_label: str, key_attr: str = "key") -> None:
        ws.append([])
        ws.append([_bold(ws, title)])
        ws.append([key_label, "成交笔数", "手数", "正负相抵 USD", "只补正数 USD"])
        for g in rows:
            ws.append([getattr(g, key_attr), g.deals, g.lots, g.comp_net_usd, g.comp_positive_usd])

    table("按账户", d.by_account, "账户")
    table("按品种", d.by_symbol, "品种")
    table("按开 / 平", d.by_entry, "开平")
    table("按日（MT 服务器日）", d.by_day, "日期")
    table("不计入的其他类别", d.not_counted_by_class, "类别", "cls")


_DETAIL_COLS: list[tuple[str, str]] = [
    ("deal_id", "成交号"), ("order_id", "订单号"), ("position_id", "仓位号"),
    ("login_sid", "账户"), ("account_group", "账户组"), ("ccy", "币种"), ("symbol", "品种"),
    ("side", "方向"), ("entry", "开平"), ("cls", "类别"), ("counted", "计入"),
    ("not_counted_reason", "不计入原因"), ("lots", "手数"),
    ("req_time_srv", "请求时间（MT）"), ("fill_time_srv", "成交时间（MT）"),
    ("req_time_utc", "请求时间 UTC"), ("fill_time_utc", "成交时间 UTC"),
    ("srv_date", "MT 服务器日"), ("delay_ms", "延迟 ms"), ("ref_price", "请求价"),
    ("fill_price", "成交价"), ("worse_px", "价差（>0 更差）"), ("outcome", "结果"),
    ("comp_usd", "补偿 USD"),
]


def _detail_sheet(ws, s: SummaryResponse, rows: Iterable[Mapping]) -> None:
    ws.freeze_panes = "A4"   # below the notice row, blank row and header
    _head(ws, s)
    ws.append([_bold(ws, label) for _, label in _DETAIL_COLS])
    for r in rows:
        out = []
        for attr, _ in _DETAIL_COLS:
            v = r[attr]
            if isinstance(v, bool):
                v = "是" if v else "否"
            elif isinstance(v, dt.date):
                v = v.isoformat()
            out.append(v)
        ws.append(out)


def build_xlsx(s: SummaryResponse, rows: Iterable[Mapping]) -> bytes:
    """``rows``: row dicts with the ``OrderRow`` field names (extra keys ignored)."""
    wb = Workbook(write_only=True)
    _summary_sheet(wb.create_sheet("汇总"), s)
    _detail_sheet(wb.create_sheet("明细"), s, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
