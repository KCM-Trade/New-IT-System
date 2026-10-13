"""Service-level tests of the CS fund risk-control ledger.

Nothing here reaches the MySQL replica or PostgreSQL:

* derivation rules (spec §5) are tested on plain row dicts through
  ``order_lookup.derive_order``;
* ``fetch_orders`` is driven through a fake pymysql connection;
* the incident service runs against ``FakePG``, a small in-memory stand-in for
  the two ledger tables that understands exactly the statements
  ``services/cs_ledger/incidents.py`` issues (and fails loudly on any other).

``FakePG`` is also imported by tests/test_audit_cs_fund_risk.py.
"""

from __future__ import annotations

import re
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal

import pymysql
import pytest

from app.core.cs_ledger_pg import HISTORY_TABLE, TABLE
from app.schemas.cs_fund_risk import IncidentCreate, IncidentUpdate
from app.services.cs_ledger import incidents, order_lookup
from app.services.cs_ledger.order_lookup import (
    OrderLookup,
    ReplicaUnavailable,
    crm_time_to_utc_iso,
    derive_order,
)
from app.services.cs_ledger.text import mask_text

ACTOR = "staff@kohleservices.com"


# ── fixtures / fakes ─────────────────────────────────────────────────────────

def replica_row(**over) -> dict:
    """One row shaped like the spec §5 SELECT."""
    base = {
        "id": 700450, "type": "deposit", "status": "approved", "pspId": 109,
        "psp_name": "CNY - Exlink CNY Wechat", "psp_cid": 0, "is_fee": 0,
        "requestedAmount": Decimal("10000.00"), "requestedCurrency": "CNY",
        "vendorAmount": Decimal("1400.00"), "vendorCurrency": "USD",
        "createdAt": datetime(2025, 5, 15, 9, 3, 46),
        "processedAt": datetime(2025, 5, 15, 9, 8, 21),
        "fromUserId": 146012, "fromLoginSid": "1-8520313",
        "firstName": "San", "middleName": None, "lastName": "Zhang",
        "is_employee": 0, "mt_group": "KCM\\STD", "mt_name": "Zhang San",
        "venue_tags": "hzl,szu",
    }
    base.update(over)
    return base


def make_lookup(transaction_id: int = 700450, **over) -> OrderLookup:
    return derive_order(replica_row(id=transaction_id, **over))


_DEFAULTS = {
    "transfer_succeeded": "unknown", "frozen_scope": "unknown", "authority": "unknown",
    "source": "manual", "version": 1, "is_test_account": False,
}
_ALL_COLUMNS = (
    "id", "incident_group_id", "transaction_id", *incidents.SNAPSHOT_FIELDS, "snapshot_at",
    *incidents.STAFF_FIELDS, "source", "import_batch", "import_sheet_row", "import_raw",
    "version", "created_by", "created_at", "updated_by", "updated_at", "deleted_at", "deleted_by",
)


class FakePG:
    """In-memory stand-in for cs_ledger.fund_risk_incidents (+ history)."""

    def __init__(self):
        self.rows: dict[int, dict] = {}
        self.history: list[dict] = []
        self.next_id = 1
        self.next_group = 100
        self.commits = 0
        self.rollbacks = 0
        self.statements: list[str] = []

    # -- connection protocol ---------------------------------------------------
    @contextmanager
    def conn(self, settings=None):
        snapshot = ({k: dict(v) for k, v in self.rows.items()}, list(self.history))
        try:
            yield self
            self.commits += 1
        except Exception:
            self.rows, self.history = snapshot   # rollback
            self.rollbacks += 1
            raise

    @contextmanager
    def cursor(self):
        yield _FakeCursor(self)

    # -- helpers ----------------------------------------------------------------
    def live(self) -> list[dict]:
        return [r for r in self.rows.values() if r["deleted_at"] is None]

    def seed(self, **over) -> dict:
        """Insert a row directly (as the import script would)."""
        now = datetime(2026, 10, 13, 1, 0, 0, tzinfo=timezone.utc)
        row = {c: None for c in _ALL_COLUMNS}
        row.update(_DEFAULTS)
        row.update(id=self.next_id, incident_group_id=self.next_group, created_by="seed",
                   updated_by="seed", created_at=now, updated_at=now, direction="deposit",
                   txn_type="deposit", channel="Exlink", txn_status="approved")
        self.next_id += 1
        self.next_group += 1
        row.update(over)
        self.rows[row["id"]] = row
        return row


class _FakeCursor:
    def __init__(self, db: FakePG):
        self.db = db
        self._result: list[dict] = []
        self.rowcount = -1

    def fetchone(self):
        return self._result.pop(0) if self._result else None

    def fetchall(self):
        out, self._result = self._result, []
        return out

    def _with_group_size(self, row: dict) -> dict:
        size = sum(1 for r in self.db.live() if r["incident_group_id"] == row["incident_group_id"])
        return {**row, "group_size": size}

    def execute(self, sql: str, params=()):
        db = self.db
        params = list(params or ())
        s = " ".join(sql.split())
        db.statements.append(s)

        if s.startswith(f"SELECT id FROM {TABLE} WHERE transaction_id = %s AND deleted_at IS NULL"):
            self._result = [{"id": r["id"]} for r in db.live() if r["transaction_id"] == params[0]]
        elif s.startswith(f"SELECT 1 AS ok FROM {TABLE} WHERE incident_group_id = %s"):
            hit = any(r["incident_group_id"] == params[0] for r in db.live())
            self._result = [{"ok": 1}] if hit else []
        elif s.startswith(f"INSERT INTO {TABLE} ("):
            m = re.match(rf"INSERT INTO {re.escape(TABLE)} \((.*?)\) VALUES \((.*?)\) ON CONFLICT", s)
            assert m, s
            assert "DO NOTHING RETURNING id" in s
            columns = [c.strip() for c in m.group(1).split(",")]
            placeholders = [p.strip() for p in m.group(2).split(",")]
            assert len(columns) == len(placeholders)
            now = datetime.now(timezone.utc)
            row = {c: None for c in _ALL_COLUMNS}
            row.update(_DEFAULTS)
            it = iter(params)
            for col, ph in zip(columns, placeholders):
                assert col in row, f"unknown column {col}"
                row[col] = next(it) if ph == "%s" else now
            assert next(it, None) is None, "too many parameters"
            if any(r["transaction_id"] == row["transaction_id"] for r in db.live()):
                self._result = []
                return
            row["id"] = db.next_id
            db.next_id += 1
            if row["incident_group_id"] is None:
                row["incident_group_id"] = db.next_group
                db.next_group += 1
            row["created_at"] = row["updated_at"] = now
            if row["amount_cny"] is not None:
                row["amount_cny"] = Decimal(str(row["amount_cny"]))
            db.rows[row["id"]] = row
            self._result = [{"id": row["id"]}]
        elif s.startswith(f"INSERT INTO {HISTORY_TABLE} "):
            incident_id, action, old_row, new_row, actor = params
            db.history.append({
                "incident_id": incident_id, "action": action, "actor": actor,
                "old_row": getattr(old_row, "adapted", old_row),
                "new_row": getattr(new_row, "adapted", new_row),
            })
            self._result = []
        elif s.startswith("SELECT t.*, (SELECT count(*)") and "WHERE t.id = %s AND t.deleted_at IS NULL" in s:
            row = db.rows.get(params[0])
            ok = row is not None and row["deleted_at"] is None
            self._result = [self._with_group_size(dict(row))] if ok else []
        elif s == f"SELECT * FROM {TABLE} WHERE id = %s AND deleted_at IS NULL FOR UPDATE":
            row = db.rows.get(params[0])
            ok = row is not None and row["deleted_at"] is None
            self._result = [dict(row)] if ok else []
        elif s.startswith(f"SELECT id FROM {TABLE} WHERE incident_group_id = (SELECT incident_group_id"):
            # Group-wide lock taken first by apply_to_group (ordered by id).
            assert s.endswith("ORDER BY id FOR UPDATE"), s
            anchor = db.rows.get(params[0])
            self._result = [] if anchor is None else [
                {"id": r["id"]} for r in sorted(db.live(), key=lambda r: r["id"])
                if r["incident_group_id"] == anchor["incident_group_id"]
            ]
        elif s.startswith(f"SELECT * FROM {TABLE} WHERE incident_group_id = %s AND id <> %s"):
            self._result = [
                dict(r) for r in sorted(db.live(), key=lambda r: r["id"])
                if r["incident_group_id"] == params[0] and r["id"] != params[1]
            ]
        elif s.startswith(f"UPDATE {TABLE} SET deleted_at = now(), deleted_by = %s"):
            actor, incident_id = params
            row = db.rows.get(incident_id)
            if row is None or row["deleted_at"] is not None:
                self._result = []
                return
            row["deleted_at"] = datetime.now(timezone.utc)
            row["deleted_by"] = actor
            self._result = [dict(row)]
        elif s.startswith(f"UPDATE {TABLE} SET ") and "version = version + 1" in s:
            head = s[len(f"UPDATE {TABLE} SET "):s.index(", version = version + 1")]
            columns = [a.split(" = ")[0].strip() for a in head.split(", ")]
            values, (actor, incident_id, version) = params[:len(columns)], params[len(columns):]
            row = db.rows.get(incident_id)
            if row is None or row["deleted_at"] is not None or row["version"] != version:
                self._result = []
                return
            for col, value in zip(columns, values):
                assert col in incidents.STAFF_FIELDS, col
                row[col] = value
            row["version"] += 1
            row["updated_by"] = actor
            row["updated_at"] = datetime.now(timezone.utc)
            self._result = [dict(row)]
        else:
            raise AssertionError(f"FakePG does not know this statement: {s}")


@pytest.fixture
def pg(monkeypatch):
    db = FakePG()
    monkeypatch.setattr(incidents, "risk_cases_conn", db.conn)
    return db


@pytest.fixture
def replica(monkeypatch):
    """Scripted replica: ``replica[txn_id] = OrderLookup``; counts calls."""
    orders: dict[int, OrderLookup] = {}
    state = {"down": False, "calls": 0}

    def fetch_orders(ids, settings=None):
        state["calls"] += 1
        if state["down"]:
            raise ReplicaUnavailable("down")
        return {i: orders[i] for i in ids if i in orders}

    def fetch_order(transaction_id, settings=None):
        return fetch_orders([transaction_id]).get(transaction_id)

    monkeypatch.setattr(order_lookup, "fetch_orders", fetch_orders)
    monkeypatch.setattr(order_lookup, "fetch_order", fetch_order)
    orders_state = type("Replica", (), {"orders": orders, "state": state})
    return orders_state


# ── derivation rules (spec §5) ────────────────────────────────────────────────

def test_deposit_takes_requested_amount_and_direction():
    look = derive_order(replica_row())
    s = look.snapshot
    assert (s.direction, s.txn_type, s.amount_cny) == ("deposit", "deposit", 10000.0)
    assert look.recordable and look.reject_reason is None and look.warnings == []


def test_withdrawal_takes_absolute_vendor_amount():
    look = derive_order(replica_row(
        type="withdrawal", requestedAmount=Decimal("2800"), requestedCurrency="USD",
        vendorAmount=Decimal("-19965.00"), vendorCurrency="CNY",
    ))
    assert look.snapshot.direction == "withdrawal"
    assert look.snapshot.amount_cny == 19965.0


def test_ib_withdrawal_is_a_withdrawal():
    look = derive_order(replica_row(type="ib withdrawal", vendorAmount=Decimal("60356"), vendorCurrency="CNY"))
    assert look.recordable
    assert (look.snapshot.direction, look.snapshot.txn_type) == ("withdrawal", "ib withdrawal")
    assert look.snapshot.amount_cny == 60356.0


@pytest.mark.parametrize("currency", ["cny", "CNY", " Cny "])
def test_currency_comparison_ignores_case(currency):
    assert derive_order(replica_row(requestedCurrency=currency)).snapshot.amount_cny == 10000.0


def test_non_cny_amount_is_none_with_a_warning():
    look = derive_order(replica_row(requestedCurrency="USD"))
    assert look.snapshot.amount_cny is None
    assert "amount_not_cny" in look.warnings
    # A deposit never falls back to the vendor side, a withdrawal never to the requested side.
    look = derive_order(replica_row(type="withdrawal", vendorCurrency="USD", requestedCurrency="CNY"))
    assert look.snapshot.amount_cny is None


def test_times_are_shifted_three_hours_to_utc():
    s = derive_order(replica_row()).snapshot
    assert s.txn_created_at == "2025-05-15T06:03:46Z"
    assert s.txn_processed_at == "2025-05-15T06:08:21Z"
    # Crossing midnight, and a fresh order with no processedAt.
    assert crm_time_to_utc_iso(datetime(2026, 1, 1, 1, 30, 0)) == "2025-12-31T22:30:00Z"
    assert crm_time_to_utc_iso("2026-01-01 01:30:00") == "2025-12-31T22:30:00Z"
    assert derive_order(replica_row(processedAt=None, status="fresh")).snapshot.txn_processed_at is None


def test_venue_is_the_first_tag_upper_cased():
    assert derive_order(replica_row(venue_tags="hzl,szu")).snapshot.venue == "HZL"
    assert derive_order(replica_row(venue_tags=None)).snapshot.venue is None


@pytest.mark.parametrize("over", [
    {"is_employee": 1},
    {"firstName": "Test", "lastName": "User"},
    {"mt_name": "DEMO account"},
    {"mt_group": "KCM\\Demo\\STD"},
    {"mt_group": "manager-TEST"},
])
def test_test_account_flag(over):
    look = derive_order(replica_row(**over))
    assert look.snapshot.is_test_account is True
    assert "test_account" in look.warnings


def test_ordinary_account_is_not_flagged():
    look = derive_order(replica_row())
    assert look.snapshot.is_test_account is False


def test_warnings_for_status_and_psp():
    look = derive_order(replica_row(status="fresh", psp_cid=1))
    assert look.warnings == ["non_cn_psp", "not_approved"]
    look = derive_order(replica_row(status="declined", requestedCurrency="USD", pspId=9999,
                                    psp_name="Brand New PSP", is_employee=1, psp_cid=2))
    assert look.warnings == ["psp_unmapped", "non_cn_psp", "test_account", "not_approved", "amount_not_cny"]


def test_channel_mapping_and_unmapped_psp():
    assert derive_order(replica_row(pspId=87)).snapshot.channel == "Npay(支付宝)"
    assert derive_order(replica_row(pspId=108)).snapshot.channel == "Npay(支付宝)"
    assert derive_order(replica_row(pspId=33)).snapshot.channel == "CC"
    assert derive_order(replica_row(pspId=6)).snapshot.channel == "Jpay"
    look = derive_order(replica_row(pspId=9999, psp_name="Brand New PSP"))
    assert look.snapshot.channel == "其他: Brand New PSP"
    assert "psp_unmapped" in look.warnings and look.recordable   # recordable, only warned
    look = derive_order(replica_row(pspId=None, psp_name=None, psp_cid=None))
    assert look.snapshot.channel.startswith("其他: ")
    assert look.warnings == ["psp_unmapped"]


def test_non_recordable_types_still_return_a_lookup():
    look = derive_order(replica_row(type="transfer in"))
    assert not look.recordable and "transfer in" in look.reject_reason
    assert look.snapshot.direction == "withdrawal"
    look = derive_order(replica_row(type="internal deposit"))
    assert not look.recordable and look.snapshot.direction == "deposit"
    look = derive_order(replica_row(is_fee=1))
    assert not look.recordable and "isFee" in look.reject_reason


def test_client_name_and_ids():
    s = derive_order(replica_row(middleName="Q")).snapshot
    assert s.client_name == "San Q Zhang"
    assert (s.client_id, s.login_sid, s.psp_id) == (146012, "1-8520313", 109)
    s = derive_order(replica_row(fromLoginSid=None, firstName=None, lastName=None)).snapshot
    assert s.login_sid is None and s.client_name is None


# ── fetch_orders against a fake pymysql connection ────────────────────────────

class _FakeMySQL:
    def __init__(self, rows, fail=None):
        self.rows, self.fail = rows, fail
        self.queries: list[tuple[str, list]] = []
        self.closed = False

    def cursor(self):
        return self

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def execute(self, sql, params=None):
        if self.fail:
            raise self.fail
        self.queries.append((sql, list(params)))
        self._out = [r for r in self.rows if r["id"] in params]

    def fetchall(self):
        return self._out

    def close(self):
        self.closed = True


def test_fetch_orders_chunks_and_omits_missing(monkeypatch):
    fake = _FakeMySQL([replica_row(id=i) for i in range(1, 1201, 2)])
    monkeypatch.setattr(order_lookup, "connect_readonly", lambda settings, **kw: fake)
    out = order_lookup.fetch_orders(list(range(1, 1201)) + [5, 5], settings=object())
    assert sorted(out) == list(range(1, 1201, 2))
    assert [len(p) for _, p in fake.queries] == [500, 500, 200]
    assert all("WHERE t.id IN (%s" in sql for sql, _ in fake.queries)
    assert fake.closed
    assert order_lookup.fetch_order(2, settings=object()) is None
    assert order_lookup.fetch_orders([], settings=object()) == {}


@pytest.mark.parametrize("exc", [
    pymysql.err.OperationalError(2003, "cannot connect"),
    pymysql.err.OperationalError(3024, "max_execution_time exceeded"),
    TimeoutError("read timed out"),
])
def test_fetch_orders_maps_connection_errors(monkeypatch, exc):
    monkeypatch.setattr(order_lookup, "connect_readonly", lambda settings, **kw: _FakeMySQL([], fail=exc))
    with pytest.raises(ReplicaUnavailable):
        order_lookup.fetch_orders([1], settings=object())

    def refuse(settings, **kw):
        raise exc

    monkeypatch.setattr(order_lookup, "connect_readonly", refuse)
    with pytest.raises(ReplicaUnavailable):
        order_lookup.fetch_orders([1], settings=object())


# ── mask_text (spec §6) ───────────────────────────────────────────────────────

def test_mask_text():
    assert mask_text(None) is None
    assert mask_text("") == ""
    assert mask_text("卡号6222021234567890123被冻结") == "卡号[卡号]被冻结"          # 19 digits
    assert mask_text("card 6222021234567890 frozen") == "card [卡号] frozen"       # 16 digits
    assert mask_text("联系电话13812345678，已核实") == "联系电话[手机号]，已核实"
    assert mask_text("13812345678") == "[手机号]"
    # Not masked: order ids, 12-15 digit runs, 20+ digit runs, 11 digits not starting with 1.
    for text in ("订单 1279140", "123456789012345", "12345678901234567890", "23812345678", "138123456789"):
        assert mask_text(text) == text
    both = mask_text("6222021234567890123 / 13812345678")
    assert both == "[卡号] / [手机号]"
    assert mask_text(both) == both


# ── incident service on FakePG ────────────────────────────────────────────────

def test_create_stores_snapshot_masks_text_and_writes_history(pg, replica):
    replica.orders[700450] = make_lookup()
    body = IncidentCreate(
        transaction_id=700450, reported_date=date(2026, 10, 13), authority="bank",
        event_text="  客户卡 6222021234567890123 被封，电话 13812345678 ", client_bank_city="  ",
        client_bank_province="广东",
    )
    result = incidents.create_incident(body, ACTOR)
    row = result.row
    assert row["transaction_id"] == 700450 and row["channel"] == "Exlink"
    assert row["client_id"] == 146012 and row["venue"] == "HZL" and row["direction"] == "deposit"
    assert row["event_text"] == "客户卡 [卡号] 被封，电话 [手机号]"
    assert row["client_bank_city"] is None            # blank -> NULL
    assert row["created_by"] == row["updated_by"] == ACTOR
    assert row["source"] == "manual" and row["version"] == 1 and row["group_size"] == 1
    assert "client_name" not in row                   # never persisted
    assert [h["action"] for h in pg.history] == ["create"]
    assert pg.history[0]["actor"] == ACTOR and pg.history[0]["old_row"] is None
    inc = incidents.to_incident(row, result.lookup)
    assert inc.order_is_live and inc.order.client_name == "San Zhang"
    assert inc.created_at.endswith("Z")


def test_create_duplicate_carries_the_existing_id(pg, replica):
    replica.orders[700450] = make_lookup()
    first = incidents.create_incident(IncidentCreate(transaction_id=700450), ACTOR).row
    with pytest.raises(incidents.DuplicateTransaction) as err:
        incidents.create_incident(IncidentCreate(transaction_id=700450), ACTOR)
    assert err.value.existing_id == first["id"]
    assert len(pg.live()) == 1 and len(pg.history) == 1


def test_a_soft_deleted_transaction_can_be_recorded_again(pg, replica):
    replica.orders[700450] = make_lookup()
    first = incidents.create_incident(IncidentCreate(transaction_id=700450), ACTOR).row
    old = incidents.soft_delete_incident(first["id"], "boss@kohleservices.com")
    assert old["deleted_at"] is None                  # the row as it stood before
    assert pg.rows[first["id"]]["deleted_by"] == "boss@kohleservices.com"
    assert pg.history[-1]["action"] == "delete"
    with pytest.raises(incidents.IncidentNotFound):
        incidents.get_row(first["id"])
    with pytest.raises(incidents.IncidentNotFound):
        incidents.soft_delete_incident(first["id"], ACTOR)
    second = incidents.create_incident(IncidentCreate(transaction_id=700450), ACTOR).row
    assert second["id"] != first["id"]


def test_create_rejects_missing_and_non_recordable_orders_before_touching_pg(pg, replica):
    with pytest.raises(incidents.OrderNotFound):
        incidents.create_incident(IncidentCreate(transaction_id=1), ACTOR)
    replica.orders[2] = make_lookup(2, type="transfer in")
    with pytest.raises(incidents.OrderNotRecordable):
        incidents.create_incident(IncidentCreate(transaction_id=2), ACTOR)
    replica.state["down"] = True
    with pytest.raises(ReplicaUnavailable):
        incidents.create_incident(IncidentCreate(transaction_id=3), ACTOR)
    assert pg.statements == []


def test_group_join_and_group_size(pg, replica):
    replica.orders[1] = make_lookup(1)
    replica.orders[2] = make_lookup(2)
    first = incidents.create_incident(IncidentCreate(transaction_id=1), ACTOR).row
    second = incidents.create_incident(
        IncidentCreate(transaction_id=2, incident_group_id=first["incident_group_id"]), ACTOR
    ).row
    assert second["incident_group_id"] == first["incident_group_id"]
    assert second["group_size"] == 2 and incidents.get_row(first["id"])["group_size"] == 2
    replica.orders[3] = make_lookup(3)
    with pytest.raises(incidents.GroupNotFound):
        incidents.create_incident(IncidentCreate(transaction_id=3, incident_group_id=987654), ACTOR)


def test_update_bumps_version_and_noop_changes_nothing(pg, replica):
    replica.orders[1] = make_lookup(1)
    row = incidents.create_incident(IncidentCreate(transaction_id=1, authority="bank"), ACTOR).row
    upd = incidents.update_incident(
        row["id"], IncidentUpdate(version=1, authority="bank", result_text="已解封 13812345678"), "b@x"
    )
    assert upd.row["version"] == 2 and upd.row["updated_by"] == "b@x"
    assert upd.row["result_text"] == "已解封 [手机号]"
    assert len(upd.changes) == 1 and upd.changes[0].target == "transaction:1"
    assert upd.changes[0].old["result_text"] is None
    assert [h["action"] for h in pg.history] == ["create", "update"]

    # Same values again: no version bump, no history row, no changes.
    again = incidents.update_incident(
        row["id"], IncidentUpdate(version=2, authority="bank", result_text="已解封 13812345678"), "c@x"
    )
    assert again.changes == [] and again.row["version"] == 2 and again.row["updated_by"] == "b@x"
    assert len(pg.history) == 2


def test_update_with_a_stale_version_conflicts(pg, replica):
    replica.orders[1] = make_lookup(1)
    row = incidents.create_incident(IncidentCreate(transaction_id=1), ACTOR).row
    incidents.update_incident(row["id"], IncidentUpdate(version=1, authority="alipay"), ACTOR)
    with pytest.raises(incidents.VersionConflict) as err:
        incidents.update_incident(row["id"], IncidentUpdate(version=1, authority="bank"), ACTOR)
    assert err.value.current_version == 2
    assert pg.rows[row["id"]]["authority"] == "alipay"
    with pytest.raises(incidents.IncidentNotFound):
        incidents.update_incident(999, IncidentUpdate(version=1), ACTOR)


def test_apply_to_group_copies_three_fields_to_siblings(pg, replica):
    for i in (1, 2, 3):
        replica.orders[i] = make_lookup(i)
    a = incidents.create_incident(IncidentCreate(transaction_id=1), ACTOR).row
    gid = a["incident_group_id"]
    b = incidents.create_incident(
        IncidentCreate(transaction_id=2, incident_group_id=gid, freeze_days=3), ACTOR).row
    other = incidents.create_incident(IncidentCreate(transaction_id=3), ACTOR).row
    upd = incidents.update_incident(
        a["id"],
        IncidentUpdate(version=1, event_text="同一事件", result_text="处理中", authority="bank",
                       freeze_days=7, apply_to_group=True),
        ACTOR,
    )
    assert [c.target for c in upd.changes] == ["transaction:1", "transaction:2"]
    sibling = pg.rows[b["id"]]
    assert (sibling["event_text"], sibling["result_text"], sibling["authority"]) == ("同一事件", "处理中", "bank")
    assert sibling["freeze_days"] == 3 and sibling["version"] == 2     # only the three fields
    assert pg.rows[other["id"]]["version"] == 1 and pg.rows[other["id"]]["event_text"] is None
    # Without the flag the sibling is untouched.
    incidents.update_incident(
        a["id"], IncidentUpdate(version=2, event_text="改了", result_text="处理中", authority="bank",
                                freeze_days=7), ACTOR)
    assert pg.rows[b["id"]]["event_text"] == "同一事件"


def test_province_is_validated_only_when_it_changes(pg, replica):
    replica.orders[1] = make_lookup(1)
    with pytest.raises(incidents.InvalidProvince):
        incidents.create_incident(IncidentCreate(transaction_id=1, id_province="广东省"), ACTOR)
    legacy = pg.seed(transaction_id=5, id_province="广东省", source="import")
    upd = incidents.update_incident(
        legacy["id"], IncidentUpdate(version=1, id_province="广东省", result_text="ok"), ACTOR)
    assert upd.row["result_text"] == "ok"
    with pytest.raises(incidents.InvalidProvince):
        incidents.update_incident(legacy["id"], IncidentUpdate(version=2, id_province="粤"), ACTOR)


def test_hydrate_uses_live_values_and_falls_back_to_the_snapshot(pg, replica):
    replica.orders[1] = make_lookup(1)
    row = incidents.create_incident(IncidentCreate(transaction_id=1), ACTOR).row
    legacy = incidents.get_row(pg.seed(transaction_id=None, direction=None, source="import")["id"])
    gone = incidents.get_row(pg.seed(transaction_id=77, amount_cny=Decimal("12.50"))["id"])

    replica.orders[1] = make_lookup(1, status="declined")       # the live value moved on
    data, live = incidents.hydrate([row, legacy, gone])
    assert live is True
    assert data[0].order_is_live and data[0].order.txn_status == "declined"
    assert data[0].order.client_name == "San Zhang"
    assert data[1].order is None and data[1].transaction_id is None and not data[1].order_is_live
    assert not data[2].order_is_live and data[2].order.amount_cny == 12.5   # missing on replica

    replica.state["down"] = True
    data, live = incidents.hydrate([row])
    assert live is False and not data[0].order_is_live
    assert data[0].order.txn_status == "approved" and data[0].order.client_name is None


def test_list_filter_sql_is_parameterised():
    f = incidents.ListFilters(
        date_from=date(2026, 1, 1), date_to=date(2026, 1, 31), channel="Jpay", direction="deposit",
        authority="bank", frozen_scope="one_card", pending_result=True, q="1279140",
    )
    where, params = incidents._where(f)
    assert params == [date(2026, 1, 1), date(2026, 1, 31), "Jpay", "deposit", "bank", "one_card",
                      1279140, 1279140]
    assert where.startswith("t.deleted_at IS NULL")
    assert "Jpay" not in where and "1279140" not in where
    assert "t.transaction_id = %s OR t.client_id = %s" in where
    assert incidents._where(incidents.ListFilters()) == ("t.deleted_at IS NULL", [])
    assert "IS NOT NULL" in incidents._where(incidents.ListFilters(pending_result=False))[0]
    assert incidents._order_by("txn_processed_at", "asc") == (
        "ORDER BY t.txn_processed_at ASC NULLS LAST, t.id ASC")
    assert incidents._order_by("drop table", "x") == "ORDER BY t.reported_date DESC NULLS LAST, t.id DESC"
    assert (incidents.total_pages(0, 50), incidents.total_pages(51, 50)) == (0, 2)


def test_export_builds_a_workbook_with_chinese_headers():
    import io

    from openpyxl import load_workbook

    from app.services.cs_ledger import export

    db = FakePG()
    row = db.seed(
        transaction_id=1279140, reported_date=date(2026, 9, 12), amount_cny=Decimal("60356.00"),
        direction="withdrawal", authority="police_antifraud", frozen_scope="all_cards",
        venue="HZL", venue_override="SZU", is_test_account=False,
        txn_processed_at=datetime(2026, 9, 11, 7, 8, 52, tzinfo=timezone.utc),
    )
    wb = load_workbook(io.BytesIO(export.build_xlsx([row, db.seed(transaction_id=None)])))
    ws = wb["台账"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3
    record = dict(zip(rows[0], rows[1]))
    assert record["订单号"] == 1279140 and record["金额（CNY）"] == 60356
    assert record["出入金"] == "出金" and record["封卡单位"] == "反诈公安"
    assert record["被封银行卡资料"] == "名下所有" and record["场地"] == "SZU"
    assert record["订单处理时间（北京时间）"] == "2026-09-11 15:08:52"
    assert record["登记日期"] == "2026-09-12" and record["测试账户"] == "否"
    assert "客户姓名" not in rows[0]
    cd = export.content_disposition(export.filename(date(2026, 10, 13)))
    assert 'filename="' in cd and "filename*=UTF-8''" in cd and cd.count(".xlsx") == 2


def test_mask_text_handles_separators_and_country_code():
    from app.services.cs_ledger.text import mask_text, strip_control, strip_ip

    assert mask_text("6222 0212 3456 7890") == "[卡号]"
    assert mask_text("6222-0212-3456-7890") == "[卡号]"
    assert mask_text("+8613812345678") == "[手机号]"
    assert mask_text("138-1234-5678") == "[手机号]"
    # A list of order ids is not a card number.
    assert mask_text("订单 700450 700484 705497") == "订单 700450 700484 705497"
    assert strip_control("a\x0bb\nc") == "ab\nc"
    assert strip_ip("175.152.117.1 成都") == "成都"
    assert strip_ip("240e:47e:50b0:f48f:d9aa:ac12:8246:5145 河南周口") == "河南周口"
    assert strip_ip("广东") == "广东"


def test_export_survives_control_characters_and_formula_text():
    import io

    import openpyxl

    from app.services.cs_ledger.export import build_xlsx

    data = build_xlsx([{"event_text": "=1+1 a\x0bb"}])
    cells = [c for c in openpyxl.load_workbook(io.BytesIO(data)).active[2] if c.value]
    assert [(c.value, c.data_type) for c in cells] == [("=1+1 ab", "s")]
